# -*- coding: utf-8 -*-
"""Shared helpers to build polished, CJK-safe .docx documents (WPS/Word compatible)."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GOLD = RGBColor(0x9A, 0x74, 0x2E)
INK = RGBColor(0x1A, 0x1C, 0x2E)
MUTED = RGBColor(0x66, 0x66, 0x66)
CJK = "Microsoft YaHei"      # renders correctly in WPS on Windows
LATIN = "Calibri"


def _set_cjk(run):
    run.font.name = LATIN
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), CJK)


def new_doc():
    doc = Document()
    # base style
    st = doc.styles['Normal']
    st.font.name = LATIN
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK)
    # narrow-ish margins
    for s in doc.sections:
        s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)
        s.top_margin = Inches(0.9); s.bottom_margin = Inches(0.9)
    return doc


def cover(doc, title, subtitle, meta_lines):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = GOLD; _set_cjk(r)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle); r2.font.size = Pt(15); r2.font.color.rgb = INK; _set_cjk(r2)
    doc.add_paragraph()
    for line in meta_lines:
        pm = doc.add_paragraph(); pm.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rm = pm.add_run(line); rm.font.size = Pt(10.5); rm.font.color.rgb = MUTED; _set_cjk(rm)
    doc.add_page_break()


def h1(doc, text):
    p = doc.add_paragraph(); p.space_before = Pt(14)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(19); r.font.color.rgb = GOLD; _set_cjk(r)
    # bottom border
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '4'); bottom.set(qn('w:color'), 'C8A15A')
    pbdr.append(bottom); pPr.append(pbdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = INK; _set_cjk(r)
    return p


def para(doc, text, size=11, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
    _set_cjk(r)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        rb = p.add_run(bold_prefix); rb.bold = True; _set_cjk(rb)
    r = p.add_run(text); _set_cjk(r)
    p.paragraph_format.line_spacing = 1.2
    return p


def numbered(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Number')
    if bold_prefix:
        rb = p.add_run(bold_prefix); rb.bold = True; _set_cjk(rb)
    r = p.add_run(text); _set_cjk(r)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True; run.font.size = Pt(10.5); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); _set_cjk(run)
        _shade(hdr[i], '9A742E')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10); _set_cjk(run)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def quote(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(11); r.font.color.rgb = MUTED; _set_cjk(r)
    p.paragraph_format.left_indent = Inches(0.3)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '10'); left.set(qn('w:color'), 'C8A15A')
    pbdr.append(left); pPr.append(pbdr)
    return p
